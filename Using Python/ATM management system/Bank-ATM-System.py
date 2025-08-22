import openpyxl
from datetime import datetime

# Account Class
class Account:
    def __init__(self, account_number, account_holder, balance):
        self.account_number = account_number
        self.account_holder = account_holder
        self.balance = balance
        self.transactions = []

    def deposit(self, amount):
        if amount > 0:
            self.balance += amount
            transaction = Transaction(
                transaction_id=f"T{len(self.transactions) + 1:03}",
                account_number=self.account_number,
                amount=amount,
                transaction_type="Deposit",
                date=datetime.now()
            )
            self.transactions.append(transaction)
            print(f"Deposit successful! New balance: ${self.balance:.2f}")
        else:
            print("Invalid deposit amount.")

    def withdraw(self, amount):
        if 0 < amount <= self.balance:
            self.balance -= amount
            transaction = Transaction(
                transaction_id=f"T{len(self.transactions) + 1:03}",
                account_number=self.account_number,
                amount=amount,
                transaction_type="Withdrawal",
                date=datetime.now()
            )
            self.transactions.append(transaction)
            print(f"Withdrawal successful! New balance: ${self.balance:.2f}")
        else:
            print("Invalid withdrawal amount or insufficient balance.")

    def check_balance(self):
        print(f"Current Balance: ${self.balance:.2f}")

    def view_transaction_history(self):
        if self.transactions:
            print("Transaction History:")
            for t in self.transactions:
                print(f"ID: {t.transaction_id}, Type: {t.transaction_type}, "
                      f"Amount: ${t.amount:.2f}, Date: {t.date.strftime('%Y-%m-%d')}")
        else:
            print("No transactions found.")


# Transaction Class
class Transaction:
    def __init__(self, transaction_id, account_number, amount, transaction_type, date):
        self.transaction_id = transaction_id
        self.account_number = account_number
        self.amount = amount
        self.transaction_type = transaction_type
        self.date = date

    def to_dict(self):
        return {
            "Transaction ID": self.transaction_id,
            "Account Number": self.account_number,
            "Amount": self.amount,
            "Transaction Type": self.transaction_type,
            "Date": self.date.strftime('%Y-%m-%d')
        }


# ATM Class
class ATM:
    def __init__(self):
        self.accounts = []

    def load_data(self):
        try:
            # Load accounts
            wb = openpyxl.load_workbook("accounts.xlsx")
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                account = Account(
                    account_number=str(row[0]).strip(),
                    account_holder=row[1],
                    balance=row[2]
                )
                self.accounts.append(account)

            # Load transactions
            wb = openpyxl.load_workbook("transactions.xlsx")
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                for account in self.accounts:
                    if account.account_number == str(row[1]).strip():
                        transaction = Transaction(
                            transaction_id=row[0],
                            account_number=str(row[1]).strip(),
                            amount=row[2],
                            transaction_type=row[3],
                            date=datetime.strptime(row[4], '%Y-%m-%d')
                        )
                        account.transactions.append(transaction)
        except FileNotFoundError:
            print("No existing data found. Starting fresh.")
        except Exception as e:
            print(f"Error loading data: {e}")

    def save_data(self):
        # Save accounts
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Account Number", "Account Holder", "Balance"])
        for acc in self.accounts:
            ws.append([acc.account_number, acc.account_holder, acc.balance])
        wb.save("accounts.xlsx")

        # Save transactions
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Transaction ID", "Account Number", "Amount", "Transaction Type", "Date"])
        for acc in self.accounts:
            for t in acc.transactions:
                ws.append([t.transaction_id, t.account_number, t.amount, t.transaction_type, t.date.strftime('%Y-%m-%d')])
        wb.save("transactions.xlsx")

    def login(self, account_number):
        # Ensure account_number is compared as a string
        account_number = account_number.strip()
        for account in self.accounts:
            if account.account_number == account_number:
                return account
        return None

    def menu(self, account):
        while True:
            print("\n1. Check Balance")
            print("2. Withdraw Money")
            print("3. Deposit Money")
            print("4. View Transaction History")
            print("5. Exit")
            choice = input("Choose an option: ")
            if choice == "1":
                account.check_balance()
            elif choice == "2":
                try:
                    amount = float(input("Enter amount to withdraw: "))
                    account.withdraw(amount)
                except ValueError:
                    print("Invalid input. Please enter a numeric value.")
            elif choice == "3":
                try:
                    amount = float(input("Enter amount to deposit: "))
                    account.deposit(amount)
                except ValueError:
                    print("Invalid input. Please enter a numeric value.")
            elif choice == "4":
                account.view_transaction_history()
            elif choice == "5":
                print("Exiting... Goodbye!")
                break
            else:
                print("Invalid option. Please try again.")


# Main Program
if __name__ == "__main__":
    atm = ATM()
    atm.load_data()

    print("Welcome to the ATM System!")
    account_number = input("Enter your account number: ")
    account = atm.login(account_number)
    if account:
        print(f"Login successful! Welcome, {account.account_holder}")
        atm.menu(account)
    else:
        print("Account not found. Exiting...")

    atm.save_data()
